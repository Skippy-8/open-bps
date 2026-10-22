"""GHG building workbooks recalculate to the engine's figures and contain no typed numbers."""
from __future__ import annotations

import re

import openpyxl
import pytest

from open_bps import DISCLAIMER, excel
from open_bps.engine import boston as boston_engine
from open_bps.engine import nyc as nyc_engine
from open_bps.inputs import BuildingRecord

formulas = pytest.importorskip("formulas")
from test_excel_cli import _solve  # noqa: E402


def _assert_checks_pass(path):
    sol = _solve(path)
    results = [v for (s, a), v in sol.items() if s == "CHECKS" and a.startswith("D") and a != "D1"]
    assert len(results) >= 6 and set(results) == {"PASS"}, results
    assert openpyxl.load_workbook(path)["Building"]["A1"].value == DISCLAIMER


def _typed_numbers(path):
    bad = []
    for ws in openpyxl.load_workbook(path).worksheets:
        for r in ws.iter_rows():
            for c in r:
                if isinstance(c.value, str) and c.value.startswith("="):
                    body = re.sub(r'"[^"]*"', "", c.value)
                    body = re.sub(r"(\w+!)?\$?[A-Z]{1,3}\$?\d+", "", body)
                    body = re.sub(r"\b(Limit_P|Use_)\w*", "", body)                 # defined names
                    if [x for x in re.findall(r"(?<![A-Za-z_])\d+(?:\.\d+)?", body) if x not in ("0", "1")]:
                        bad.append((ws.title, c.coordinate, c.value))
    return bad


def test_nyc_mixed_use_workbook(tmp_path):
    b = BuildingRecord(building_id="ex", name="DOB example", gross_floor_area=45_000,
                       spaces=[{"property_type": "Multifamily Housing", "gfa": 40_000},
                               {"property_type": "Personal Services (Health/Beauty, Dry Cleaning, etc)", "gfa": 1_900},
                               {"property_type": "Vehicle Repair Services", "gfa": 600},
                               {"property_type": "Retail Store", "gfa": 2_500}],
                       utilities={"electricity_kwh": 900_000, "natural_gas_kbtu": 2_250_000, "fuel_oil_2_kbtu": 1_050_000})
    path = excel.building_workbook("nyc_ll97", b, tmp_path / "nyc.xlsx", 2024, 2050)
    _assert_checks_pass(path)
    assert not _typed_numbers(path)
    ex = openpyxl.load_workbook(path)["Exposure"]
    assert ex.cell(2, 1).value == 2024 and ex.cell(28, 1).value == 2050


def test_nyc_workbook_from_the_scored_dataset(nyc_scored, tmp_path):
    row = nyc_scored[nyc_scored["screen_status"].eq("scored") & (nyc_scored["penalty_2030"] > 0)].iloc[0].to_dict()
    path = excel.building_workbook("nyc_ll97", nyc_engine.record(row), tmp_path / "n.xlsx", 2024, 2034)
    _assert_checks_pass(path)


def test_boston_workbook_with_rps_renewables_and_steam_override(boston_scored, tmp_path):
    s = boston_scored[boston_scored["screen_status"].eq("scored") & boston_scored["acp_2025"].gt(0)]
    row = s[s["district_steam_factor_t_per_kbtu"].notna()].iloc[0].to_dict() if s["district_steam_factor_t_per_kbtu"].notna().any() \
        else s.iloc[0].to_dict()
    path = excel.building_workbook("boston_berdo", boston_engine.record(row), tmp_path / "b.xlsx", 2025, 2050,
                                   overrides=boston_engine.overrides(row))
    _assert_checks_pass(path)
    assert not _typed_numbers(path)
    sol = _solve(path)
    acp_2025 = sol[("EXPOSURE", f"{openpyxl.utils.get_column_letter(openpyxl.load_workbook(path)['Exposure'].max_column - 1)}2")]
    assert acp_2025 == pytest.approx(row["acp_2025"], abs=0.01)


def test_workbook_refuses_an_unpriceable_first_year(tmp_path):
    b = BuildingRecord(building_id="x", gross_floor_area=1000, spaces=[{"property_type": "Office", "gfa": 1000}],
                       utilities={"electricity_kwh": 1, "fuel_oil_5_6_kbtu": 10})
    with pytest.raises(ValueError, match="fuel_oil_5_6"):
        excel.building_workbook("nyc_ll97", b, tmp_path / "x.xlsx", 2024, 2030)


def test_cli_nyc_and_boston(root, built, nyc_scored, boston_scored, capsys, tmp_path):
    from open_bps import cli
    for city, df in (("nyc", nyc_scored), ("boston", boston_scored)):
        bid = df[df["screen_status"].eq("scored")]["building_id"].iloc[0]
        cli.main(["--root", str(root), "lookup", city, bid])
        assert bid in capsys.readouterr().out
        out = tmp_path / f"{city}.xlsx"
        cli.main(["--root", str(root), "excel", city, bid, "--out", str(out)])
        assert out.exists()
    cli.main(["--root", str(root), "validate"])
    assert "verdict_match_rate" in capsys.readouterr().out
    cli.main(["rules"])
    assert "nyc: Local Law 97" in capsys.readouterr().out
