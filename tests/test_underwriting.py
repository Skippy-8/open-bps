"""Retrofit capital stack: §179D, financing, NPV/IRR/payback, the grid toggle, and workbook parity."""
from __future__ import annotations

from datetime import date

import openpyxl
import pytest

from open_bps import excel
from open_bps.inputs import BuildingRecord
from open_bps.underwriting.capital_stack import Financing, deduction_179d, underwrite, underwrite_building
from open_bps.underwriting.finance import fv, irr, npv, pmt

BEFORE, AFTER = date(2026, 3, 1), date(2026, 7, 1)


# ---- §179D -------------------------------------------------------------------------------------------
def test_179d_rates_follow_the_irs_tables():
    # Rev. Proc. 2025-32 §4.25: 2026 base $0.59 + $0.02/point, cap $1.19; PWA $2.97 + $0.12/point, cap $5.94
    assert deduction_179d(100_000, 0.25, 2026, BEFORE).per_sqft == pytest.approx(0.59)
    assert deduction_179d(100_000, 0.40, 2026, BEFORE).per_sqft == pytest.approx(0.59 + 15 * 0.02)
    assert deduction_179d(100_000, 0.70, 2026, BEFORE).per_sqft == pytest.approx(1.19)          # capped
    assert deduction_179d(100_000, 0.30, 2026, BEFORE, pwa=True).per_sqft == pytest.approx(2.97 + 5 * 0.12)
    assert deduction_179d(100_000, 0.90, 2026, BEFORE, pwa=True).per_sqft == pytest.approx(5.94)
    # Form 7205 instructions: 2025 base 0.58 (cap 1.16), PWA 2.90 + 0.12 (cap 5.81); 2024 PWA 2.83 + 0.11
    assert deduction_179d(1, 0.54, 2025, BEFORE).per_sqft == pytest.approx(1.16)
    assert deduction_179d(1, 0.35, 2024, BEFORE, pwa=True).per_sqft == pytest.approx(2.83 + 10 * 0.11)
    # the worksheet rounds the percentage to two decimals first: 26.4% -> 0.26
    assert deduction_179d(1, 0.264, 2026, BEFORE).per_sqft == pytest.approx(0.61)


def test_179d_gates_and_limits():
    assert not deduction_179d(100_000, 0.4, 2026, None).eligible                               # no start date
    late = deduction_179d(100_000, 0.4, 2027, AFTER)
    assert not late.eligible and late.amount == 0 and "70507" in late.note
    assert deduction_179d(100_000, 0.4, 2026, date(2026, 6, 30)).eligible                       # the last day
    assert not deduction_179d(100_000, 0.24, 2026, BEFORE).eligible                             # below 25%
    d = deduction_179d(100_000, 0.4, 2026, BEFORE, tax_rate=0.21, cost_basis=50_000)
    assert d.amount == 50_000 and d.cash_value == pytest.approx(10_500) and "cost" in d.note
    later = deduction_179d(100_000, 0.4, 2028, BEFORE)
    assert later.per_sqft == pytest.approx(0.89) and "only through 2026" in later.note


# ---- finance -----------------------------------------------------------------------------------------
def test_amortisation_matches_closed_form():
    p = pmt(0.06, 20, -1_000_000)
    assert p == pytest.approx(1_000_000 * 0.06 / (1 - 1.06 ** -20))                            # 87,184.56
    bal10 = fv(0.06, 10, p, -1_000_000)
    assert bal10 == pytest.approx(p * (1 - 1.06 ** -10) / 0.06)                                 # PV of the last 10
    assert fv(0.06, 20, p, -1_000_000) == pytest.approx(0, abs=1e-6)


def test_npv_irr_payback_by_hand():
    u = underwrite([2026, 2027, 2028, 2029], [300, 400, 500, 200], [0, 0, 0, 0], 1000, Financing.cash(), 0.08)
    assert u.irr_unlevered == pytest.approx(0.1532213788, abs=1e-9)                             # Excel IRR
    assert u.npv_unlevered == pytest.approx(-1000 + 300 / 1.08 + 400 / 1.08 ** 2 + 500 / 1.08 ** 3 + 200 / 1.08 ** 4)
    assert u.payback_year == 2028 and u.equity == 1000 and u.npv_levered == pytest.approx(u.npv_unlevered)
    never = underwrite([2026], [10], [0], 1000, Financing.cash(), 0.08)
    assert never.payback_year is None


def test_financing_flows():
    years, base = list(range(2026, 2036)), [200_000] * 10
    cp = underwrite(years, base, [0] * 10, 1_000_000, Financing.cpace(0.065, 20), 0.08)
    pay = pmt(0.065, 20, -1_000_000)
    assert cp.equity == 0 and cp.irr_levered != cp.irr_levered                                 # n/m: no equity
    assert cp.table["debt_service"].tolist() == pytest.approx([pay] * 10)
    assert cp.table["balance_repaid"].iloc[-1] == pytest.approx(fv(0.065, 10, pay, -1_000_000))
    dt = underwrite(years, base, [0] * 10, 1_000_000, Financing.debt(0.6, 0.07, 5), 0.08)
    assert dt.equity == pytest.approx(400_000) and dt.table["debt_service"].iloc[5:].eq(0).all()
    assert dt.table["balance_repaid"].sum() == 0
    assert dt.npv_levered == pytest.approx(-400_000 + npv(0.08, dt.table["levered_flow"]))
    assert dt.irr_levered == pytest.approx(irr([-400_000, *dt.table["levered_flow"]]))
    with pytest.raises(ValueError):
        underwrite(years, base, [float("nan")] * 10, 1_000_000, Financing.cash(), 0.08)


# ---- buildings ---------------------------------------------------------------------------------------
def _office(**use):
    return BuildingRecord(building_id="o", name="Office", gross_floor_area=200_000,
                          spaces=[{"property_type": "Office", "gfa": 200_000}], utilities=use)


def test_retrofit_avoids_penalties_and_179d_is_gated():
    b = _office(electricity_kwh=4_000_000, natural_gas_kbtu=12_000_000)
    kw = dict(capex=3_000_000, eui_reduction=0.4, financing=Financing.cash(), start_year=2026, hold_years=10)
    late = underwrite_building("nyc_ll97", b, construction_start=AFTER, **kw)
    early = underwrite_building("nyc_ll97", b, construction_start=BEFORE, **kw)
    assert late.incentive == 0 and early.incentive == pytest.approx(200_000 * (0.59 + 15 * 0.02) * 0.21)
    assert (late.table["avoided_penalty"] >= 0).all() and late.table["avoided_penalty"].sum() > 0
    assert early.npv_unlevered - late.npv_unlevered == pytest.approx(early.incentive / 1.08)
    assert any("§179D" in n and "terminated" in n for n in late.notes)


def test_grid_toggle_moves_only_electricity_driven_exposure():
    kw = dict(capex=1_000_000, eui_reduction=0.3, financing=Financing.cash(), start_year=2026, hold_years=10)
    gas = _office(electricity_kwh=0, natural_gas_kbtu=30_000_000)
    s = underwrite_building("nyc_ll97", gas, **kw)
    d = underwrite_building("nyc_ll97", gas, grid="dynamic_grid_projection", grid_decline=0.05, **kw)
    assert s.table["baseline_penalty"].tolist() == d.table["baseline_penalty"].tolist()
    elec = _office(electricity_kwh=12_000_000, natural_gas_kbtu=0)
    s = underwrite_building("boston_berdo", elec, **kw)
    d = underwrite_building("boston_berdo", elec, grid="dynamic_grid_projection", grid_decline=0.05, **kw)
    assert s.table["baseline_penalty"].iloc[0] == pytest.approx(d.table["baseline_penalty"].iloc[0])
    assert d.table["baseline_penalty"].iloc[-1] < s.table["baseline_penalty"].iloc[-1]
    assert any("grid scenario" in n for n in d.notes)
    with pytest.raises(ValueError, match="energy standard"):
        underwrite_building("denver_energize", elec, **kw)


# ---- workbook parity ---------------------------------------------------------------------------------
@pytest.mark.parametrize("financing, grid", [
    (Financing.cpace(0.065, 20), "static_statutory"),           # balance repaid at the end of the hold
    (Financing.debt(0.6, 0.07, 7), "dynamic_grid_projection"),
    (Financing.cash(), "static_statutory"),
])
def test_underwriting_workbook_reconciles(financing, grid, tmp_path):
    pytest.importorskip("formulas")
    from test_excel_ghg import _assert_checks_pass, _typed_numbers

    b = BuildingRecord(building_id="m", name="Mixed", gross_floor_area=150_000,
                       spaces=[{"property_type": "Office", "gfa": 120_000}, {"property_type": "Retail Store", "gfa": 30_000}],
                       utilities={"electricity_kwh": 3_500_000, "natural_gas_kbtu": 9_000_000, "fuel_oil_2_kbtu": 400_000})
    u = underwrite_building("nyc_ll97", b, capex=2_500_000, eui_reduction=0.35, financing=financing, start_year=2026,
                            hold_years=12, energy_cost_savings=150_000, grid=grid, grid_decline=0.03,
                            construction_start=BEFORE, pwa=True)
    path = excel.building_workbook("nyc_ll97", b, tmp_path / "u.xlsx", 0, 0, underwriting=u)
    _assert_checks_pass(path)
    assert not _typed_numbers(path)
    wb = openpyxl.load_workbook(path)
    assert wb.sheetnames[:2] == ["Building", "Underwriting"]
    labels = [wb["Checks"].cell(i, 1).value for i in range(2, wb["Checks"].max_row)]
    assert "IRR unlevered" in labels and "NPV levered ($)" in labels
