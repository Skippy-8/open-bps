"""PortfolioAnalyzer and `open-bps analyze`."""
from __future__ import annotations

from pathlib import Path

import openpyxl
import pytest

from open_bps import PortfolioAnalyzer, cli
from open_bps.underwriting.capital_stack import Financing

FIX = Path(__file__).parent / "fixtures"
PORTFOLIO = """building_id,name,gross_floor_area,property_type_1,gfa_1,property_type_2,gfa_2,electricity_kwh,natural_gas_therms,fuel_oil_2_kbtu
A,Alpha,120000,Office,90000,Retail Store,30000,2600000,60000,
B,Beta,80000,Multifamily Housing,80000,,,700000,70000,300000
C,Gamma,50000,Other,50000,,,500000,10000,
"""


@pytest.fixture()
def csv(tmp_path):
    p = tmp_path / "portfolio.csv"
    p.write_text(PORTFOLIO)
    return p


def test_analyze_and_underwrite_a_portfolio(csv, tmp_path):
    pa = PortfolioAnalyzer("nyc")
    pf = pa.load(csv)
    ex = pa.analyze(pf, last=2040)
    assert list(ex["status"]) == ["ok", "ok", "cannot_price"]                       # 'Other' cannot be assigned
    assert "may not be assigned" in ex.loc[2, "problems"]
    assert ex.loc[0, "first_year_over_limit"] in (2024, 2030) and ex["rule_version"].nunique() == 1
    deals = pa.underwrite(pf, eui_reduction=0.3, financing=Financing.cpace(0.065, 20), capex_per_sqft=20,
                          start_year=2026, energy_cost_savings_per_sqft=0.5)
    assert list(deals["status"]) == ["ok", "ok", "cannot_price"]
    assert deals.loc[0, "avoided_penalty_hold"] > 0
    path = pa.to_excel(tmp_path / "model.xlsx")
    from test_excel_ghg import _assert_checks_pass, _typed_numbers
    _assert_checks_pass(path)
    assert not _typed_numbers(path)
    wb = openpyxl.load_workbook(path)
    assert wb.sheetnames[0] == "Portfolio" and "Bldg1_Underwriting" in wb.sheetnames and "Bldg3_Building" not in wb.sheetnames


def test_analyzer_refuses_energy_standards():
    with pytest.raises(ValueError, match="sets targets per building"):
        PortfolioAnalyzer("denver")


def test_espm_input(tmp_path):
    pa = PortfolioAnalyzer("nyc_ll97")
    with pytest.raises(Exception, match="row"):
        pa.load(FIX / "nyc_ll84.csv", espm=True)                                     # the real file has bad rows
    from open_bps.adapters.espm import read_espm
    r = read_espm(FIX / "nyc_ll84.csv", "nyc_ll97")
    assert len(r.buildings) > 380 and r.errors


def test_cli_analyze(csv, tmp_path, capsys):
    out = tmp_path / "m.xlsx"
    cli.main(["analyze", str(csv), "--jurisdiction", "boston", "--eui-reduction", "0.25", "--capex-per-sqft", "15",
              "--financing", "debt", "--share", "0.6", "--start-year", "2026", "--output", str(out)])
    text = capsys.readouterr().out
    assert "npv_unlevered" in text and out.exists()
    bad = tmp_path / "bad.csv"
    bad.write_text(PORTFOLIO.replace("2600000", "").replace("Retail Store", "Retial Store"))
    with pytest.raises(SystemExit) as e:
        cli.main(["analyze", str(bad), "--jurisdiction", "nyc"])
    assert e.value.code == 2
    msg = capsys.readouterr().out
    assert "Row 2 (Building 'Alpha')" in msg and "'Retial Store' is not one this jurisdiction maps" in msg
