"""One-building workbooks recalculate to the engine's values; the CLI runs offline."""
from __future__ import annotations

import re

import openpyxl
import pytest

from open_bps import DISCLAIMER, cli, excel

formulas = pytest.importorskip("formulas")


def _solve(path):
    out = {}
    for k, v in formulas.ExcelModel().loads(str(path)).finish().calculate().items():
        m = re.match(r"'\[[^\]]+\]([^']+)'!([A-Z]+\d+)$", k)
        if m:
            out[(m.group(1), m.group(2))] = v.value[0][0] if hasattr(v, "value") else v
    return out


@pytest.mark.parametrize("city", ["denver", "dc"])
def test_workbook_checks_pass(city, denver_scored, dc_scored, tmp_path):
    df = denver_scored if city == "denver" else dc_scored
    col = "extended_2032_penalty" if city == "denver" else "payment_est_performance"
    row = df[df["screen_status"].eq("scored") & (df[col] > 0)].iloc[0].to_dict()
    path = (excel.denver_workbook if city == "denver" else excel.dc_workbook)(row, tmp_path / f"{city}.xlsx")
    sol = _solve(path)
    results = [v for (s, a), v in sol.items() if s == "CHECKS" and a.startswith("D") and a != "D1"]
    assert len(results) >= 3 and set(results) == {"PASS"}
    assert openpyxl.load_workbook(path)["Building"]["A1"].value == DISCLAIMER


@pytest.mark.parametrize("city", ["denver", "dc"])
def test_no_typed_numbers_in_formulas(city, denver_scored, dc_scored, tmp_path):
    df = denver_scored if city == "denver" else dc_scored
    row = df[df["screen_status"].eq("scored")].iloc[0].to_dict()
    path = (excel.denver_workbook if city == "denver" else excel.dc_workbook)(row, tmp_path / f"{city}.xlsx")
    bad = []
    for ws in openpyxl.load_workbook(path).worksheets:
        for r in ws.iter_rows():
            for c in r:
                if isinstance(c.value, str) and c.value.startswith("="):
                    body = re.sub(r'"[^"]*"', "", c.value)
                    body = re.sub(r"(\w+!)?\$?[A-Z]{1,3}\$?\d+", "", body)        # cell references
                    lits = [x for x in re.findall(r"(?<![A-Za-z_])\d+(?:\.\d+)?", body) if x not in ("0", "1")]
                    if lits:
                        bad.append((ws.title, c.coordinate, c.value))
    assert not bad, bad[:3]


def test_cli_lookup_and_excel(root, built, denver_scored, capsys, tmp_path):
    bid = denver_scored[denver_scored["screen_status"].eq("scored")]["building_id"].iloc[0]
    cli.main(["--root", str(root), "lookup", "denver", bid])
    assert bid in capsys.readouterr().out
    out = tmp_path / "b.xlsx"
    cli.main(["--root", str(root), "excel", "denver", bid, "--out", str(out)])
    assert out.exists()
    with pytest.raises(SystemExit, match="no denver building"):
        cli.main(["--root", str(root), "excel", "denver", "does-not-exist"])
    cli.main(["--root", str(root), "validate"])
    assert "verdict_match_rate" in capsys.readouterr().out
