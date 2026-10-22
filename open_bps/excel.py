"""One-building Excel workbook: inputs with sources, live-formula exposure, and a Checks tab.

Every calculation is a formula over named input cells (no typed numbers in calculations other than 0 and
1), so a reader can audit or change any assumption. The engine's own results sit on the Checks tab as
labelled reference values; a check passes when the workbook's formulas reproduce them.
"""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.workbook.defined_name import DefinedName

from open_bps import DISCLAIMER
from open_bps import rules as _rules

BLUE, BOLD, GREY = Font(color="0000FF"), Font(bold=True), Font(color="808080", italic=True)
WARN = PatternFill("solid", fgColor="FFF2CC")
TOL = 0.01


def _name(wb: Workbook, name: str, sheet: str, cell: str) -> None:
    wb.defined_names[name] = DefinedName(name, attr_text=f"'{sheet}'!${cell[0]}${cell[1:]}")


def _head(ws, title: str, attribution: str) -> int:
    ws["A1"] = DISCLAIMER
    ws["A1"].font, ws["A1"].fill = BOLD, WARN
    ws["A2"] = title
    ws["A2"].font = Font(bold=True, size=13)
    ws["A3"] = f"Data: {attribution}"
    ws["A3"].font = GREY
    ws.column_dimensions["A"].width = 46
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 70
    return 5


def _inputs(wb, ws, start: int, items: list[tuple[str, str, object, str, str]]) -> int:
    """items: (defined name, label, value, number format, source). Returns the next free row."""
    for j, h in enumerate(["Input", "Value", "Source"]):
        ws.cell(start, 1 + j, h).font = BOLD
    r = start
    for name, label, value, fmt, source in items:
        r += 1
        ws.cell(r, 1, label)
        c = ws.cell(r, 2, value)
        c.font, c.number_format = BLUE, fmt
        ws.cell(r, 3, source).font = GREY
        _name(wb, name, ws.title, f"B{r}")
    return r + 2


def _checks(wb, rows: list[tuple]) -> None:
    """rows: (label, formula giving the workbook value, engine value[, tolerance for this row])."""
    ws = wb.create_sheet("Checks")
    for j, h in enumerate(["Check", "Workbook", "Engine (reference)", "Result"]):
        ws.cell(1, 1 + j, h).font = BOLD
    _name(wb, "Tolerance", "Checks", f"B{len(rows) + 4}")
    for i, (label, formula, engine_value, *tol) in enumerate(rows, start=2):
        ws.cell(i, 1, label)
        ws.cell(i, 2, formula).number_format = "#,##0.00"
        c = ws.cell(i, 3, engine_value)
        c.font, c.number_format = BLUE, "#,##0.00"
        if tol:
            t = ws.cell(i, 5, tol[0])
            t.font, t.number_format = BLUE, "0.0E+00"
            for col in (2, 3):
                ws.cell(i, col).number_format = "0.000000%"
            ws.cell(i, 4, f'=IF(ABS(B{i}-C{i})<=E{i},"PASS","FAIL")').font = BOLD
        else:
            ws.cell(i, 4, f'=IF(ABS(B{i}-C{i})<=Tolerance,"PASS","FAIL")').font = BOLD
    n = len(rows) + 2
    ws.cell(n, 1, "All checks pass").font = BOLD
    ws.cell(n, 4, f'=IF(COUNTIF(D2:D{n - 1},"PASS")=ROWS(D2:D{n - 1}),"PASS","FAIL")').font = BOLD
    ws.cell(n + 2, 1, "Tolerance ($ or kBtu)")
    ws.cell(n + 2, 2, TOL).font = BLUE
    if any(len(r) > 3 for r in rows):
        ws.cell(1, 5, "Row tolerance").font = BOLD
    ws.column_dimensions["A"].width = 44
    for col in "BC":
        ws.column_dimensions[col].width = 18


def _sources(wb, r: dict) -> None:
    ws = wb.create_sheet("Sources")
    ws["A1"], ws["A1"].font = f"Rules: {r['rule_version']} (checked {r['checked_on']})", BOLD
    for i, (key, src) in enumerate(r["sources"].items(), start=3):
        ws.cell(i, 1, key)
        ws.cell(i, 2, src.get("title", src.get("section", "")))
        ws.cell(i, 3, src["url"])
    ws.column_dimensions["B"].width = 70


def denver_workbook(row: dict, out: Path, through: int = 2035) -> Path:
    """`row`: one scored Denver building (canonical + engine columns)."""
    r = _rules.load("denver")
    wb = Workbook()
    ws = wb.active
    ws.title = "Building"
    start = _head(ws, f"Energize Denver exposure: {row['property_name']} (Building ID {row['building_id']})",
                  "Energize Denver Master Dataset (City and County of Denver), as is; not for engineering purposes")
    tg = "Energize Denver Technical Guidance v3 (April 2025)"
    ext, orig = r["schedules"]["extended"], r["schedules"]["original"]
    nxt = _inputs(wb, ws, start, [
        ("GFA", "Gross floor area (sq ft)", row["gfa"], "#,##0", "Master Dataset: Master Sq Ft"),
        ("EUI", "Weather-normalized site EUI (kBtu/sq ft), held flat", row["eui"], "0.0",
         f"Master Dataset, data year {int(row['data_year']) if row.get('data_year') == row.get('data_year') else 'n/a'}; before city credits, so an upper bound"),
        ("T_First", "First interim target EUI", row["target_first_interim"], "0.0", "Master Dataset"),
        ("T_Second", "Second interim target EUI", row["target_second_interim"], "0.0", "Master Dataset"),
        ("T_Final", "Adjusted final target EUI", row["target_final"], "0.0", "Master Dataset"),
        ("Rate_Ext", "Penalty rate, extended schedule ($/kBtu)", ext["rate_per_kbtu"], "$0.00", f"{tg}, Table 17 (two periods)"),
        ("Rate_Orig", "Penalty rate, original schedule ($/kBtu)", orig["rate_per_kbtu"], "$0.00", f"{tg}, Table 17 (three periods)"),
    ])
    ws.cell(nxt, 1, "Exposure if EUI stays at its latest value").font = BOLD
    hdr = nxt + 1
    for j, h in enumerate(["Schedule", "Year", "Target EUI", "kBtu over target", "Penalty ($)"]):
        ws.cell(hdr, 1 + j, h).font = BOLD
    lines = [("Extended (2028/2032)", ext["target_years"]["first_interim"], "T_First", "Rate_Ext"),
             ("Extended (2028/2032)", ext["target_years"]["final"], "T_Final", "Rate_Ext")]
    lines += [("Extended, annual after final", y, "T_Final", "Rate_Ext") for y in range(ext["target_years"]["final"] + 1, through + 1)]
    lines += [("Original (2025/2027/2030)", orig["target_years"][k], t, "Rate_Orig")
              for k, t in (("first_interim", "T_First"), ("second_interim", "T_Second"), ("final", "T_Final"))]
    first = hdr + 1
    for i, (label, year, target, rate) in enumerate(lines, start=first):
        ws.cell(i, 1, label)
        ws.cell(i, 2, year)
        ws.cell(i, 3, f"={target}").number_format = "0.0"
        ws.cell(i, 4, f"=MAX(0,(EUI-C{i})*GFA)").number_format = "#,##0"
        ws.cell(i, 5, f"=D{i}*{rate}").number_format = "$#,##0"
    n_ext = 2 + (through - ext["target_years"]["final"])
    last_ext = first + n_ext - 1
    tot = first + len(lines) + 1
    ws.cell(tot, 1, f"Extended schedule, cumulative through {through}").font = BOLD
    ws.cell(tot, 5, f"=SUM(E{first}:E{last_ext})").number_format = "$#,##0"
    ws.cell(tot + 1, 1, "Reduction needed to reach the final target").font = BOLD
    ws.cell(tot + 1, 5, "=MAX(0,1-T_Final/EUI)").number_format = "0.0%"
    _name(wb, "Cum_Ext", "Building", f"E{tot}")

    yrs = ext["target_years"]
    _checks(wb, [
        (f"Penalty {yrs['first_interim']} (extended)", f"=Building!E{first}", row[f"extended_{yrs['first_interim']}_penalty"]),
        (f"Penalty {yrs['final']} (extended)", f"=Building!E{first + 1}", row[f"extended_{yrs['final']}_penalty"]),
        (f"Cumulative through {through} (extended)", "=Cum_Ext", row[f"extended_cumulative_penalty_through_{through}"]),
        (f"Penalty {orig['target_years']['final']} (original)", f"=Building!E{first + len(lines) - 1}",
         row[f"original_{orig['target_years']['final']}_penalty"]),
    ])
    _sources(wb, r)
    out = Path(out)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out


def dc_workbook(row: dict, out: Path) -> Path:
    """`row`: one scored DC building."""
    r = _rules.load("dc")
    wb = Workbook()
    ws = wb.active
    ws.title = "Building"
    start = _head(ws, f"DC BEPS Cycle 1 exposure: {row['property_name']} (Portfolio Manager ID {row['building_id']})",
                  "DOEE via Open Data DC (CC BY 4.0)")
    gb = "DOEE BEPS Guidebook v1.2 (Aug 2026)"
    num = lambda v: None if v != v else v  # noqa: E731  (NaN -> blank cell)
    nxt = _inputs(wb, ws, start, [
        ("GFA", "Gross floor area (sq ft)", row["gfa"], "#,##0", "DOEE disclosure (reported GFA)"),
        ("Metric_Type", "BEPS metric type", row["metric_type"], "@", "DOEE disclosure"),
        ("Standard", "BEPS standard for the property type", num(row["standard"]), "0.0", "DOEE disclosure"),
        ("Metric", "Building's BEPS metric (metric year)", num(row["metric"]), "0.0", "DOEE disclosure (frozen baseline)"),
        ("Base_EUI", "Site EUI, average of 2018-2019", num(row["site_eui_baseline"]), "0.0", "DOEE benchmarking; Estimated proxy for Adjusted Site EUI"),
        ("Latest_EUI", "Site EUI, latest reported year", num(row["site_eui_latest"]), "0.0", "DOEE benchmarking"),
        ("Req_Cut", "Performance Pathway required site EUI reduction", r["pathways"]["performance"]["required_reduction"], "0%", f"{gb}, Ch. 3.1"),
        ("Pay_PSF", "Maximum payment per sq ft ($)", r["max_payment"]["per_sq_ft_gfa"], "$0.00", f"{gb}, Ch. 7"),
        ("Pay_Cap", "Maximum payment cap per building ($)", r["max_payment"]["cap_per_building"], "$#,##0", f"{gb}, Ch. 7"),
    ])
    ws.cell(nxt, 1, "Results").font = BOLD
    rows = [("Meets the BEPS standard (metric year)", '=IF(Metric_Type="ENERGYSTARSCORE",Metric>=Standard,Metric<=Standard)', "General"),
            ("Maximum alternative compliance payment", "=MIN(GFA*Pay_PSF,Pay_Cap)", "$#,##0"),
            ("Site EUI reduction so far (Estimated)", "=IF(Base_EUI>0,1-Latest_EUI/Base_EUI,0)", "0.0%"),
            ("Performance Pathway payment if latest year held (Estimated)",
             "=MIN(GFA*Pay_PSF,Pay_Cap)*(1-MIN(1,MAX(0,IF(Base_EUI>0,1-Latest_EUI/Base_EUI,0))/Req_Cut))", "$#,##0")]
    for i, (label, f, fmt) in enumerate(rows, start=nxt + 1):
        ws.cell(i, 1, label)
        ws.cell(i, 2, f).number_format = fmt
    _name(wb, "Max_Pay", "Building", f"B{nxt + 2}")
    _name(wb, "PP_Pay", "Building", f"B{nxt + 4}")
    checks = [("Maximum payment", "=Max_Pay", row["max_payment"])]
    if row.get("payment_est_performance") == row.get("payment_est_performance"):
        checks.append(("Performance Pathway payment (Estimated)", "=PP_Pay", row["payment_est_performance"]))
    _checks(wb, checks)
    _sources(wb, r)
    out = Path(out)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out



# ---- GHG jurisdictions (NYC LL97, Boston BERDO) -------------------------------------------------------------
def _col(i: int) -> str:
    s = ""
    while i:
        i, r = divmod(i - 1, 26)
        s = chr(65 + r) + s
    return s


def building_workbook(jurisdiction: str, building, out: Path, first: int, last: int, blend_standard: bool = False,
                      overrides: dict | None = None, data_note: str = "your data", underwriting=None) -> Path:
    """A GHG building's exposure year by year, every figure a formula over named inputs.

    Spaces: each space's sq ft times its class's limit per compliance period, summed into the building limit.
    Exposure: per year, fuel use times that year's factors (electricity after the zero-emission share and any
    renewable kWh), the limit for the year's period, the excess and the penalty. Checks: the engine's figures.
    Years run from `first` to `last`, stopping before any year the policy cannot price (e.g. no factor yet).

    With `underwriting` (from open_bps.underwriting.capital_stack.underwrite_building) the years are the hold
    period, Exposure adds the post-retrofit columns, and an Underwriting sheet builds the capital stack.
    """
    from open_bps.registry import get_policy

    wb = Workbook()
    wb.active.title = "Building"
    checks, _ = _ghg_sheets(wb, "", jurisdiction, building, first, last, blend_standard, overrides, data_note, underwriting)
    _checks(wb, checks)
    _ghg_sources(wb, get_policy(jurisdiction))
    out = Path(out)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out


def _ghg_sources(wb, p) -> None:
    ws_s = wb.create_sheet("Sources")
    ws_s["A1"], ws_s["A1"].font = f"Rules: {p.rule_version} (checked {p.checked_on})", BOLD
    for i, (key, s_) in enumerate(p.sources.items(), start=3):
        ws_s.cell(i, 1, key)
        ws_s.cell(i, 2, s_.title or s_.section or "")
        ws_s.cell(i, 3, s_.url)
    ws_s.column_dimensions["B"].width = 70


def _ghg_sheets(wb, tag: str, jurisdiction: str, building, first: int, last: int, blend_standard: bool,
                overrides: dict | None, data_note: str, underwriting) -> tuple[list[tuple], dict]:
    """Write one building's sheets into `wb`, prefixing sheet titles and defined names with `tag` (empty for a
    single-building workbook). Returns the Checks rows and references to its key results."""
    from open_bps.engine import blend, ghg
    from open_bps.policy import FACTOR_UNITS
    from open_bps.registry import get_policy

    p = get_policy(jurisdiction)
    u = underwriting
    ui = u.inputs if u is not None else {}
    if u is not None:
        first, last = int(u.table["year"].iloc[0]), int(u.table["year"].iloc[-1])
    cut = ui.get("eui_reduction")

    def year_overrides(y: int) -> dict:
        ov = dict(overrides or {})
        if ui.get("grid") == "dynamic_grid_projection" and "electricity" not in ov:
            ov["electricity"] = p.factor("electricity", first)[0] * (1 - ui["grid_decline"]) ** (y - first)
        return ov

    use = {f: v for f, v in building.utilities.by_fuel().items() if v}
    results, post = [], []
    for y in range(first, last + 1):
        r = ghg.assess(p, building, y, blend_standard=blend_standard, overrides=year_overrides(y))
        if r.problems:
            if not results or u is not None:
                raise ValueError(f"cannot price {y}: " + "; ".join(r.problems))
            break
        results.append(r)
        if cut is not None:
            post.append(ghg.assess(p, building, y, blend_standard=blend_standard, overrides=year_overrides(y),
                                   use={f: v * (1 - cut) for f, v in building.utilities.by_fuel().items()}))
    years = [r.year for r in results]
    periods = [q for q in p.compliance_periods if any(q.covers(y) for y in years)]

    N = lambda n: tag + n  # noqa: E731   defined names, unique per building
    ref = lambda sheet: f"'{tag}{sheet}'!"  # noqa: E731
    ws = wb[f"{tag}Building"]
    start = _head(ws, f"{p.program} exposure: {building.name or building.building_id} (ID {building.building_id})", data_note)
    src = lambda key: p.sources[key].title or key  # noqa: E731
    renew_applies = bool(p.electricity and p.electricity.subtract_renewable_kwh)
    items = [(N("GFA"), "Gross floor area (sq ft)", building.gross_floor_area, "#,##0", data_note),
             (N("Rate"), "Penalty per tCO2e over the limit ($)", ghg.penalty_rate(p), "$#,##0",
              f"{src(p.penalty.source)}; {p.penalty.section or ''}"),
             (N("Renew"), "Renewable electricity deducted (kWh)",
              building.utilities.renewable_electricity_kwh if renew_applies else 0, "#,##0",
              data_note if renew_applies else f"{p.id} has no such deduction")]
    for f, v in use.items():
        items.append((N(f"Use_{f}"), f"{f.replace('_', ' ').capitalize()} ({'kWh' if f == 'electricity' else 'kBtu'} a year), held flat",
                      v, "#,##0", data_note))
    if cut is not None:
        items.append((N("Cut"), "Retrofit: cut in every fuel's use", cut, "0.0%", "retrofit assumption (user input)"))
    _inputs(wb, ws, start, items)

    # Spaces: limits by period for the class each space is counted at
    sp = wb.create_sheet(f"{tag}Spaces")
    lim0 = blend.limit(p, list(building.spaces), years[0], blend=blend_standard)
    head = ["Property type", "Law's class", "Counted at", "Sq ft"] + [f"Limit {q.id} (t/sf)" for q in periods] \
        + [f"Contribution {q.id} (t)" for q in periods]
    for j, h in enumerate(head, start=1):
        sp.cell(1, j, h).font = BOLD
    n = len(lim0.spaces)
    for i, s_ in enumerate(lim0.spaces, start=2):
        sp.cell(i, 1, s_.property_type)
        sp.cell(i, 2, s_.statutory_class)
        sp.cell(i, 3, s_.applied_class)
        c = sp.cell(i, 4, s_.gfa)
        c.font, c.number_format = BLUE, "#,##0"
        for k, q in enumerate(periods):
            c = sp.cell(i, 5 + k, p.threshold(s_.applied_class, q.start_year))
            c.font, c.number_format = BLUE, "0.000000000"
            sp.cell(i, 5 + len(periods) + k, f"=D{i}*{_col(5 + k)}{i}").number_format = "#,##0.00"
    tot = n + 3
    sp.cell(tot, 1, "Building limit (tCO2e a year)").font = BOLD
    for k in range(len(periods)):
        cc = _col(5 + len(periods) + k)
        sp.cell(tot, 5 + len(periods) + k, f"=SUM({cc}2:{cc}{n + 1})").number_format = "#,##0.00"
        _name(wb, N(f"Limit_P{k + 1}"), f"{tag}Spaces", f"{cc}{tot}")
    th = p.thresholds
    sp.cell(tot + 2, 1, f"Limits: {src(th.source)}, {th.section or ''} (converted from {th.unit} to t/sf)").font = GREY
    sp.cell(tot + 3, 1, f"Method: {lim0.method}. " + " ".join(lim0.notes)).font = GREY
    sp.column_dimensions["A"].width = 44

    # Exposure: one row per year
    ex = wb.create_sheet(f"{tag}Exposure")
    fuels = list(use)
    head = ["Year", "Period", "Zero-emission share of grid kWh"] \
        + [f"Factor: {f} (t per {'kWh' if f == 'electricity' else 'kBtu'})" for f in fuels] \
        + ["Electricity counted (kWh)", "Emissions (tCO2e)", "Limit (tCO2e)", "Excess (tCO2e)", "Penalty ($)", "Estimated"]
    if cut is not None:
        head += ["Retrofit: electricity counted (kWh)", "Retrofit: emissions (tCO2e)", "Retrofit: excess (tCO2e)",
                 "Retrofit: penalty ($)"]
    for j, h in enumerate(head, start=1):
        ex.cell(1, j, h).font = BOLD
    fc = {f: _col(4 + j) for j, f in enumerate(fuels)}
    nf = len(fuels)
    c_cnt, c_em, c_lim, c_exc, c_pen = (_col(4 + nf + k) for k in range(5))
    c_cnt2, c_em2, c_exc2, c_pen2 = (_col(10 + nf + k) for k in range(4))
    for i, r in enumerate(results, start=2):
        y = r.year
        ov = year_overrides(y)
        ex.cell(i, 1, y)
        ex.cell(i, 2, r.period)
        c = ex.cell(i, 3, p.rps_share(y))
        c.font, c.number_format = BLUE, "0%"
        terms, terms2 = [], []
        for f in fuels:
            c = ex.cell(i, 4 + fuels.index(f), ov[f] if f in ov else p.factor(f, y)[0])
            c.font, c.number_format = BLUE, "0.000000000"
            if f == "electricity":
                terms.append(f"{c_cnt}{i}*{fc[f]}{i}")
                terms2.append(f"{c_cnt2}{i}*{fc[f]}{i}")
            else:
                terms.append(f"{N('Use_' + f)}*{fc[f]}{i}")
                terms2.append(f"{N('Use_' + f)}*(1-{N('Cut')})*{fc[f]}{i}")
        has_e = "electricity" in use
        ex.cell(i, 4 + nf, f"=MAX(0,{N('Use_electricity')}*(1-C{i})-{N('Renew')})" if has_e else 0)
        ex.cell(i, 5 + nf, "=" + ("+".join(terms) or "0")).number_format = "#,##0.00"
        k = next(j for j, q in enumerate(periods) if q.covers(y))
        ex.cell(i, 6 + nf, f"={N('Limit_P')}{k + 1}").number_format = "#,##0.00"
        ex.cell(i, 7 + nf, f"=MAX(0,{c_em}{i}-{c_lim}{i})").number_format = "#,##0.00"
        ex.cell(i, 8 + nf, f"={c_exc}{i}*{N('Rate')}").number_format = "$#,##0"
        ex.cell(i, 9 + nf, "Estimated" if r.estimated else "")
        if cut is not None:
            ex.cell(i, 10 + nf, f"=MAX(0,{N('Use_electricity')}*(1-{N('Cut')})*(1-C{i})-{N('Renew')})" if has_e else 0)
            ex.cell(i, 11 + nf, "=" + ("+".join(terms2) or "0")).number_format = "#,##0.00"
            ex.cell(i, 12 + nf, f"=MAX(0,{c_em2}{i}-{c_lim}{i})").number_format = "#,##0.00"
            ex.cell(i, 13 + nf, f"={c_exc2}{i}*{N('Rate')}").number_format = "$#,##0"
    last_row = len(results) + 1
    ex.cell(last_row + 2, 1, f"Cumulative penalty {years[0]}-{years[-1]}").font = BOLD
    ex.cell(last_row + 2, 8 + nf, f"=SUM({c_pen}2:{c_pen}{last_row})").number_format = "$#,##0"
    _name(wb, N("Cum_Penalty"), f"{tag}Exposure", f"{c_pen}{last_row + 2}")
    notes = [f"{f}: {src(p.emission_factors[f].source)}, {p.emission_factors[f].section or ''} "
             f"(converted from {p.emission_factors[f].unit}; {FACTOR_UNITS[p.emission_factors[f].unit][0]} basis)"
             for f in fuels if f in p.emission_factors and f not in (overrides or {})]
    notes += [f"{f}: factor supplied by the user" for f in (overrides or {}) if f in use]
    if ui.get("grid") == "dynamic_grid_projection":
        notes.append(f"electricity: grid scenario, the {first} factor declining {ui['grid_decline']:.1%} a year (not law)")
    if p.electricity:
        notes.append(f"Zero-emission share: {src(p.electricity.source)}, {p.electricity.section or ''}")
    notes.append("Use is held at the latest values ('cost of inaction'). 'Estimated' rows use factors from guidance, not law.")
    for j, t in enumerate(notes):
        ex.cell(last_row + 4 + j, 1, t).font = GREY
    ex.column_dimensions["A"].width = 10

    lbl = f"{building.building_id}: " if tag else ""
    pick = sorted({0, len(results) // 2, len(results) - 1})
    checks = [(lbl + f"Limit {results[0].year} (t)", f"={ref('Exposure')}{c_lim}2", results[0].limit_t)]
    for j in pick:
        r = results[j]
        checks.append((lbl + f"Emissions {r.year} (t)", f"={ref('Exposure')}{c_em}{j + 2}", r.emissions_t))
        checks.append((lbl + f"Penalty {r.year} ($)", f"={ref('Exposure')}{c_pen}{j + 2}", r.penalty))
    checks.append((lbl + f"Cumulative penalty {years[0]}-{years[-1]} ($)", f"={N('Cum_Penalty')}", sum(r.penalty for r in results)))
    if u is not None:
        for j in pick:
            checks.append((lbl + f"Retrofit penalty {post[j].year} ($)", f"={ref('Exposure')}{c_pen2}{j + 2}", post[j].penalty))
        uw_checks, uw_refs = _underwriting_sheet(wb, tag, u, f"{ref('Exposure')}{c_pen}", f"{ref('Exposure')}{c_pen2}")
        checks += uw_checks
    refs = {"cum_penalty": N("Cum_Penalty"), "first_year": years[0], "last_year": years[-1]}
    if u is not None:
        refs.update(uw_refs)
    return checks, refs


def _underwriting_sheet(wb, tag: str, u, base_col: str, post_col: str) -> tuple[list[tuple], dict]:
    """The capital stack as formulas over the Exposure sheet's penalties. Returns the Checks rows and references."""
    def N(n: str) -> str:            # defined names, unique per building
        return tag + n
    ws = wb.create_sheet(f"{tag}Underwriting", wb.sheetnames.index(f"{tag}Building") + 1)
    sheet = f"'{tag}Underwriting'!"
    lbl = f"{u.inputs.get('building_id')}: " if tag else ""
    f = u.financing
    ws["A1"], ws["A1"].font = "Retrofit underwriting (screening estimate; inputs in blue)", Font(bold=True, size=13)
    ws.column_dimensions["A"].width = 46
    nxt = _inputs(wb, ws, 3, [
        (N("Capex"), "Retrofit capex ($), paid at t = 0", u.capex, "$#,##0", "user input"),
        (N("Fin_Share"), f"Share of capex financed ({f.kind})", f.share if f.kind != "cash" else 0, "0%", "user input"),
        (N("Int_Rate"), "Financing rate (annual)", f.rate, "0.00%", "user input"),
        (N("Term"), "Financing term (years)", f.term_years, "0", "user input"),
        (N("Disc_Rate"), "Discount rate", u.discount_rate, "0.00%", "user input"),
        (N("Savings"), "Energy cost savings ($ a year)", float(u.table["energy_savings"].iloc[0]), "$#,##0",
         "user input (0 = not given)"),
        (N("Incentive"), "§179D deduction x tax rate, received in year 1 ($)", u.incentive, "$#,##0",
         next((n for n in u.notes if n.startswith("§179D")), "§179D")),
    ])
    capex, share, rate, term, disc = N("Capex"), N("Fin_Share"), N("Int_Rate"), N("Term"), N("Disc_Rate")
    savings, incentive, fin, pay, equity = N("Savings"), N("Incentive"), N("Fin_Amt"), N("Payment"), N("Equity")
    r0 = nxt
    for label, formula, fmt in [("Amount financed", f"={capex}*{share}", "$#,##0"),
                                ("Annual debt service", f"=IF({fin}=0,0,PMT({rate},{term},-{fin}))", "$#,##0.00"),
                                ("Equity at t = 0", f"={capex}-{fin}", "$#,##0")]:
        ws.cell(nxt, 1, label)
        ws.cell(nxt, 2, formula).number_format = fmt
        nxt += 1
    _name(wb, fin, f"{tag}Underwriting", f"B{r0}")
    _name(wb, pay, f"{tag}Underwriting", f"B{r0 + 1}")
    _name(wb, equity, f"{tag}Underwriting", f"B{r0 + 2}")

    hdr = nxt + 1
    cols = ["Year", "t", "Baseline penalty", "Retrofit penalty", "Avoided penalty", "Energy savings", "§179D",
            "Unlevered flow", "Debt service", "Balance repaid", "Levered flow", "Cumulative unlevered", "Paid back in"]
    for j, h in enumerate(cols, start=1):
        ws.cell(hdr, j, h).font = BOLD
    t0 = hdr + 1
    ws.cell(t0, 1, "t = 0")
    ws.cell(t0, 2, 0)
    ws.cell(t0, 8, f"=-{capex}").number_format = "$#,##0"
    ws.cell(t0, 11, f"=-{equity}").number_format = "$#,##0"
    ws.cell(t0, 12, f"=H{t0}").number_format = "$#,##0"
    h = len(u.table)
    for k in range(h):
        i = t0 + 1 + k
        ex_row = k + 2
        ws.cell(i, 1, f"='{tag}Exposure'!A{ex_row}")
        ws.cell(i, 2, f"=B{i - 1}+1")
        ws.cell(i, 3, f"={base_col}{ex_row}").number_format = "$#,##0"
        ws.cell(i, 4, f"={post_col}{ex_row}").number_format = "$#,##0"
        ws.cell(i, 5, f"=C{i}-D{i}").number_format = "$#,##0"
        ws.cell(i, 6, f"={savings}").number_format = "$#,##0"
        ws.cell(i, 7, f"=IF(B{i}=1,{incentive},0)").number_format = "$#,##0"
        ws.cell(i, 8, f"=E{i}+F{i}+G{i}").number_format = "$#,##0"
        ws.cell(i, 9, f"=IF(B{i}<={term},{pay},0)").number_format = "$#,##0"
        ws.cell(i, 10, f"=IF({term}>B{i},FV({rate},B{i},{pay},-{fin}),0)" if k == h - 1 else 0).number_format = "$#,##0"
        ws.cell(i, 11, f"=H{i}-I{i}-J{i}").number_format = "$#,##0"
        ws.cell(i, 12, f"=L{i - 1}+H{i}").number_format = "$#,##0"
        ws.cell(i, 13, f'=IF(AND(L{i}>=0,L{i - 1}<0),A{i},"")')
    end = t0 + h
    s = end + 2
    summary = [
        ("PV of avoided penalties, savings and §179D", f"=NPV({disc},H{t0 + 1}:H{end})", "$#,##0"),
        ("NPV, unlevered", f"=H{t0}+NPV({disc},H{t0 + 1}:H{end})", "$#,##0"),
        ("IRR, unlevered", f'=IFERROR(IRR(H{t0}:H{end}),"n/m")', "0.00%"),
        ("NPV, levered (equity)", f"=K{t0}+NPV({disc},K{t0 + 1}:K{end})", "$#,##0"),
        ("IRR, levered", f'=IF({equity}>0,IFERROR(IRR(K{t0}:K{end}),"n/m"),"n/m")', "0.00%"),
        ("Net capex after §179D", f"={capex}-{incentive}", "$#,##0"),
        ("Simple payback (year)", f'=IF(COUNT(M{t0 + 1}:M{end})=0,"not within the hold",MIN(M{t0 + 1}:M{end}))', "0"),
    ]
    for j, (label, formula, fmt) in enumerate(summary):
        ws.cell(s + j, 1, label).font = BOLD
        ws.cell(s + j, 2, formula).number_format = fmt
    for j, note in enumerate(u.notes):
        ws.cell(s + len(summary) + 1 + j, 1, note).font = GREY
    checks = [(lbl + "PV of benefits ($)", f"={sheet}B{s}", u.npv_benefits),
              (lbl + "NPV unlevered ($)", f"={sheet}B{s + 1}", u.npv_unlevered),
              (lbl + "NPV levered ($)", f"={sheet}B{s + 3}", u.npv_levered),
              (lbl + "Net capex after §179D ($)", f"={sheet}B{s + 5}", u.net_capex)]
    if u.irr_unlevered == u.irr_unlevered:
        checks.append((lbl + "IRR unlevered", f"={sheet}B{s + 2}", u.irr_unlevered, 1e-6))
    if u.irr_levered == u.irr_levered:
        checks.append((lbl + "IRR levered", f"={sheet}B{s + 4}", u.irr_levered, 1e-6))
    if u.payback_year is not None:
        checks.append((lbl + "Payback year", f"={sheet}B{s + 6}", u.payback_year))
    for k in (0, h - 1):
        checks.append((lbl + f"Levered flow, year {k + 1} ($)", f"={sheet}K{t0 + 1 + k}",
                       float(u.table["levered_flow"].iloc[k])))
    refs = {"npv_unlevered": f"{sheet}B{s + 1}", "irr_unlevered": f"{sheet}B{s + 2}", "payback": f"{sheet}B{s + 6}"}
    return checks, refs
