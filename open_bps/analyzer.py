"""Analyse your own buildings: exposure under a GHG standard, retrofit underwriting, and an audited workbook.

    from open_bps import PortfolioAnalyzer
    from open_bps.underwriting.capital_stack import Financing

    pa = PortfolioAnalyzer("nyc_ll97")
    portfolio = pa.load("portfolio.csv")                 # or pa.load("espm_export.xlsx", espm=True)
    exposure = pa.analyze(portfolio)                     # one row per building
    deals = pa.underwrite(portfolio, capex_per_sqft=25, eui_reduction=0.3,
                          financing=Financing.cpace(rate=0.065, term_years=20))
    pa.to_excel("model.xlsx")                            # every figure a formula, with a Checks sheet

Every result carries the rule version; every figure is a screening estimate (see open_bps.DISCLAIMER).
"""
from __future__ import annotations

from datetime import date
from pathlib import Path

import pandas as pd
from openpyxl import Workbook

from open_bps import DISCLAIMER
from open_bps.engine import ghg
from open_bps.inputs import PortfolioInput
from open_bps.registry import get_policy
from open_bps.underwriting.capital_stack import Financing, Underwriting, underwrite_building


class PortfolioAnalyzer:
    def __init__(self, jurisdiction: str, blend_standard: bool = False):
        self.policy = get_policy(jurisdiction)
        if "ghg_intensity" not in self.policy.metrics:
            raise ValueError(f"{self.policy.id} sets targets per building in its own data; use `open-bps lookup` "
                             "for it. PortfolioAnalyzer covers the GHG standards: "
                             + ", ".join(j for j in _ghg_jurisdictions()))
        self.blend_standard = blend_standard
        self.portfolio: PortfolioInput | None = None
        self.deals: dict[str, Underwriting] = {}
        self._window: tuple[int, int] | None = None

    @property
    def allowed_property_types(self) -> set[str]:
        return set(self.policy.property_type_map) | set(self.policy.unassignable_types)

    def load(self, path, espm: bool = False) -> PortfolioInput:
        """Read and validate a portfolio file (generic template, or an ESPM export with `espm=True`)."""
        if espm:
            from open_bps.adapters.espm import from_espm
            return from_espm(path, self.policy.id)
        return PortfolioInput.from_csv(path, self.allowed_property_types)

    def analyze(self, portfolio: PortfolioInput, first: int | None = None, last: int = 2050,
                hold_years: int = 10) -> pd.DataFrame:
        """Exposure per building with use held flat: penalties at the start of each compliance period, the first
        year over the limit, and the cumulative penalty over `hold_years` from `first`."""
        p = self.policy
        first = first or p.compliance_periods[0].start_year
        self.portfolio, self._window = portfolio, (first, last)
        starts = [q.start_year for q in p.compliance_periods if first <= q.start_year <= last]
        rows = []
        for b in portfolio.buildings:
            row = {"building_id": b.building_id, "name": b.name, "gross_floor_area": b.gross_floor_area,
                   "property_types": "; ".join(f"{s.property_type}: {s.gfa:g}" for s in b.spaces)}
            res = {y: ghg.assess(p, b, y, blend_standard=self.blend_standard) for y in range(first, last + 1)}
            for y in sorted({first, *starts}):
                r = res[y]
                row[f"emissions_t_{y}"], row[f"limit_t_{y}"], row[f"penalty_{y}"] = r.emissions_t, r.limit_t, r.penalty
            over = [y for y, r in res.items() if r.penalty]
            row["first_year_over_limit"] = over[0] if over else None
            hold = [res[y].penalty for y in range(first, min(last, first + hold_years - 1) + 1)]
            row[f"cumulative_penalty_{hold_years}y"] = None if any(v is None for v in hold) else sum(hold)
            problems = sorted({x for r in res.values() for x in r.problems})
            row["problems"] = "; ".join(problems) or None
            row["estimated_from"] = next((y for y, r in res.items() if r.estimated), None)
            row["status"] = "ok" if res[first].penalty is not None else "cannot_price"
            rows.append(row)
        out = pd.DataFrame(rows)
        out["rule_version"] = p.rule_version
        return out

    def underwrite(self, portfolio: PortfolioInput, eui_reduction: float, financing: Financing,
                   capex: dict | None = None, capex_per_sqft: float | None = None, start_year: int | None = None,
                   hold_years: int = 10, discount_rate: float = 0.08, energy_cost_savings_per_sqft: float = 0.0,
                   grid: str = "static_statutory", grid_decline: float = 0.0, construction_start: date | None = None,
                   pwa: bool = False, tax_rate: float = 0.21) -> pd.DataFrame:
        """Underwrite the same retrofit across the portfolio. Capex is given per building (`capex`, by id) or per
        sq ft. Buildings the standard cannot price are listed with the reason, not dropped."""
        if (capex is None) == (capex_per_sqft is None):
            raise ValueError("give capex (per building) or capex_per_sqft, not both")
        start_year = start_year or max(date.today().year, self.policy.compliance_periods[0].start_year)
        self.portfolio, self.deals = portfolio, {}
        rows = []
        for b in portfolio.buildings:
            c = capex[b.building_id] if capex is not None else capex_per_sqft * b.gross_floor_area
            row = {"building_id": b.building_id, "name": b.name, "gross_floor_area": b.gross_floor_area, "capex": c}
            try:
                u = underwrite_building(self.policy.id, b, c, eui_reduction, financing, start_year, hold_years,
                                        discount_rate, energy_cost_savings_per_sqft * b.gross_floor_area, grid,
                                        grid_decline, construction_start, pwa, tax_rate, self.blend_standard)
            except ValueError as e:
                rows.append({**row, "status": "cannot_price", "problems": str(e)})
                continue
            self.deals[b.building_id] = u
            t = u.table
            rows.append({**row, "status": "ok", "baseline_penalty_hold": float(t["baseline_penalty"].sum()),
                         "avoided_penalty_hold": float(t["avoided_penalty"].sum()), "incentive_179d": u.incentive,
                         "net_capex": u.net_capex, "npv_unlevered": u.npv_unlevered, "irr_unlevered": u.irr_unlevered,
                         "npv_levered": u.npv_levered, "irr_levered": u.irr_levered, "payback_year": u.payback_year,
                         "notes": " | ".join(u.notes)})
        out = pd.DataFrame(rows)
        out["rule_version"] = self.policy.rule_version
        out["financing"] = financing.kind
        return out

    def to_excel(self, path, last: int = 2050) -> Path:
        """One workbook for the portfolio: a summary sheet linked by formula to each building's Spaces, Exposure
        (and Underwriting) sheets, and one Checks sheet reconciling every building to the engine."""
        from open_bps import excel

        if self.portfolio is None:
            raise ValueError("run analyze() or underwrite() first")
        p = self.policy
        wb = Workbook()
        summary = wb.active
        summary.title = "Portfolio"
        summary["A1"], summary["A1"].font = DISCLAIMER, excel.BOLD
        summary["A1"].fill = excel.WARN
        summary["A2"], summary["A2"].font = f"{p.program}: portfolio ({p.rule_version})", excel.Font(bold=True, size=13)
        head = ["Building", "Name", "Sq ft", "First year", "Last year", "Cumulative penalty, use unchanged ($)"]
        if self.deals:
            head += ["NPV, unlevered ($)", "IRR, unlevered", "Payback year"]
        for j, h in enumerate(head, start=1):
            summary.cell(4, j, h).font = excel.BOLD
        checks, r = [], 5
        first = self._window[0] if self._window else None
        for k, b in enumerate(self.portfolio.buildings, start=1):
            u = self.deals.get(b.building_id)
            if self.deals and u is None:
                continue                                 # could not be priced; see underwrite()'s table
            tag = f"Bldg{k}_"
            wb.create_sheet(f"{tag}Building")
            try:
                rows, refs = excel._ghg_sheets(wb, tag, p.id, b, first or p.compliance_periods[0].start_year,
                                               self._window[1] if self._window else last, self.blend_standard,
                                               None, "your data", u)
            except ValueError as e:
                for name in [n for n in wb.sheetnames if n.startswith(tag)]:
                    del wb[name]
                summary.cell(r, 1, b.building_id)
                summary.cell(r, 2, f"not priced: {e}")
                r += 1
                continue
            checks += rows
            summary.cell(r, 1, b.building_id)
            summary.cell(r, 2, b.name)
            summary.cell(r, 3, f"={tag}GFA").number_format = "#,##0"
            summary.cell(r, 4, refs["first_year"])
            summary.cell(r, 5, refs["last_year"])
            summary.cell(r, 6, f"={refs['cum_penalty']}").number_format = "$#,##0"
            if u is not None:
                summary.cell(r, 7, f"={refs['npv_unlevered']}").number_format = "$#,##0"
                summary.cell(r, 8, f"={refs['irr_unlevered']}").number_format = "0.00%"
                summary.cell(r, 9, f"={refs['payback']}")
            r += 1
        if not checks:
            raise ValueError("no building in the portfolio could be priced")
        summary.cell(r + 1, 1, "Total").font = excel.BOLD
        summary.cell(r + 1, 6, f"=SUM(F5:F{r - 1})").number_format = "$#,##0"
        if self.deals:
            summary.cell(r + 1, 7, f"=SUM(G5:G{r - 1})").number_format = "$#,##0"
        summary.column_dimensions["A"].width = 18
        summary.column_dimensions["B"].width = 36
        excel._checks(wb, checks)
        excel._ghg_sources(wb, p)
        path = Path(path)
        path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(path)
        return path


def _ghg_jurisdictions() -> list[str]:
    from open_bps.registry import list_jurisdictions
    return [j for j in list_jurisdictions() if "ghg_intensity" in get_policy(j).metrics]


__all__ = ["PortfolioAnalyzer", "Financing"]
